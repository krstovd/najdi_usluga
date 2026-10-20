from pathlib import Path
import re, html
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

root = Path(__file__).resolve().parent
source = root / 'Najdi-Usluga-Dokumentacija.md'
lines = source.read_text(encoding='utf-8').splitlines()
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(2)
sec.left_margin = sec.right_margin = Cm(2.1)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.12
for name, size in [('Title', 30), ('Heading 1', 18), ('Heading 2', 13)]:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(16)
    s.paragraph_format.space_after = Pt(8)
    s.paragraph_format.keep_with_next = True
footer = sec.footer.paragraphs[0]
footer.alignment = 2
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
doc.core_properties.title = 'Најди Услуга документација на апликацијата'
doc.core_properties.subject = 'Функционалности архитектура и подготовка за презентација'
doc.core_properties.author = ''
parts, toc = [], []
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1; continue
    if line.startswith('```'):
        code = []; i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            code.append(lines[i]); i += 1
        p = doc.add_paragraph('\n'.join(code)); p.paragraph_format.keep_together = True
        for r in p.runs: r.font.name = 'Consolas'; r.font.size = Pt(9)
        parts.append('<pre><code>' + html.escape('\n'.join(code)) + '</code></pre>')
    elif line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+', x) for x in cells): rows.append(cells)
            i += 1
        i -= 1
        table = doc.add_table(rows=0, cols=len(rows[0])); table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = [5.1, 11.7] if len(rows[0]) == 2 else [4.0, 5.2, 7.6]
        for col, width in zip(table.columns, widths): col.width = Cm(width)
        for ri, data in enumerate(rows):
            cells = table.add_row().cells
            for ci, txt in enumerate(data):
                cell = cells[ci]; cell.width = Cm(widths[ci]); cell.text = txt
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    e = OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
                pr.append(borders)
                margin = OxmlElement('w:tcMar')
                for edge in ['top','left','bottom','right']:
                    e = OxmlElement('w:'+edge); e.set(qn('w:w'),'95'); e.set(qn('w:type'),'dxa'); margin.append(e)
                pr.append(margin)
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7EEF4' if ri == 0 else 'FFFFFF'); pr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.space_before = Pt(3)
                    for r in p.runs: r.font.size = Pt(9); r.bold = ri == 0
            if ri == 0:
                repeat = OxmlElement('w:tblHeader'); table.rows[ri]._tr.get_or_add_trPr().append(repeat)
        parts.append('<div class="table-wrap"><table><thead><tr>' + ''.join('<th>'+html.escape(c)+'</th>' for c in rows[0]) + '</tr></thead><tbody>' + ''.join('<tr>'+''.join('<td>'+html.escape(c)+'</td>' for c in row)+'</tr>' for row in rows[1:]) + '</tbody></table></div>')
    elif line.startswith('#'):
        level = len(line) - len(line.lstrip('#')); title = line[level:].strip()
        doc.add_paragraph(title, 'Title' if level == 1 else 'Heading 1' if level == 2 else 'Heading 2')
        ident = 'section-' + str(len(toc))
        if level == 2: toc.append((ident, title))
        parts.append(f'<h{level} id="{ident}">{html.escape(title)}</h{level}>')
    elif line.startswith('- '):
        doc.add_paragraph(line[2:], 'List Bullet'); parts.append('<p class="bullet">• '+html.escape(line[2:])+'</p>')
    else:
        p = doc.add_paragraph(line)
        parts.append('<p>'+html.escape(line)+'</p>')
    i += 1
work = root.parent / 'target' / 'documentation'
work.mkdir(parents=True, exist_ok=True)
doc.save(work / 'Najdi-Usluga-Dokumentacija.docx')
css = '''*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;color:#18222c;background:#f0f3f6;font:16px/1.65 "Segoe UI",Arial,sans-serif}main{max-width:980px;margin:36px auto;padding:60px 70px;background:white;box-shadow:0 8px 30px #17293b10}h1,h2,h3{color:#000;line-height:1.25;break-after:avoid}h1{font-size:46px;margin:0 0 20px}h2{font-size:26px;margin:48px 0 18px}h3{font-size:20px;margin:28px 0 12px}p{margin:0 0 14px}table{border-collapse:collapse;width:100%;font-size:14px;line-height:1.5;margin:18px 0 24px;table-layout:fixed}th,td{border:1px solid #d9d9d9;padding:11px 13px;text-align:left;vertical-align:middle;overflow-wrap:anywhere}th{background:#e7eef4;color:#000;font-weight:650}thead{display:table-header-group}tr{break-inside:avoid}td:first-child{font-weight:550}pre{padding:18px;background:#f2f5f8;border-radius:6px;white-space:pre-wrap;overflow-wrap:anywhere;font-size:14px}.bullet{padding-left:20px;text-indent:-15px}nav{padding:22px;background:#f4f7fa;margin:30px 0}nav a{display:block;color:#284c69;text-decoration:none;font-size:14px;margin:5px 0}header{max-width:980px;margin:20px auto -16px;text-align:right}button{padding:10px 20px;background:#203e56;color:white;border:0;border-radius:5px;cursor:pointer;font:inherit}@media(max-width:700px){main{margin:0;padding:30px 22px}h1{font-size:36px}table{font-size:12px}th,td{padding:7px}}@media print{@page{size:A4;margin:18mm}body{background:white;font-size:10.5pt;line-height:1.45}main{max-width:none;margin:0;padding:0;box-shadow:none}header,nav{display:none}h1{font-size:30pt}h2{font-size:18pt;margin-top:24pt}h3{font-size:13pt}table{font-size:9pt}th,td{padding:6pt}p{orphans:3;widows:3}pre{font-size:9pt}}'''
navigation = '<nav aria-label="Содржина">' + ''.join(f'<a href="#{ident}">{html.escape(title)}</a>' for ident,title in toc if re.match(r'^\d+ ',title)) + '</nav>'
parts.insert(7, navigation)
out = '<!doctype html><html lang="mk"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Најди Услуга документација</title><style>'+css+'</style></head><body><header><button onclick="window.print()">Печати или зачувај како PDF</button></header><main>'+''.join(parts)+'</main></body></html>'
(root / 'Najdi-Usluga-Dokumentacija.html').write_text(out, encoding='utf-8')
print('Created HTML and DOCX; source words:', len(source.read_text(encoding='utf-8').split()))
